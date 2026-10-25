"""웹 화면을 통한 분기 결산 전체 흐름."""

import io
from datetime import date

import openpyxl

from .conftest import xlsx


def _period_id(client, name):
    from app import db
    from app.models import Period

    with db.SessionLocal() as s:
        return s.query(Period).filter_by(name=name).one().id


def test_login_required(client):
    assert client.get("/login", follow_redirects=False).headers["location"] == "/setup"
    r = client.get("/receipts", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "/login"


def test_quarter_close_flow(admin, legacy_file):
    c = admin
    r = c.post("/periods/import-legacy", files={"file": ("old.xlsx", legacy_file)})
    assert "기준 기간을 가져왔습니다" in r.text
    assert "품목 3건" in r.text  # 'NP-0499 / 0' 자리표시 품목은 건너뜀

    r = c.post("/periods/next")
    assert "2026.3Q" in r.text and "2026-07-01 ~ 2026-09-30" in r.text
    pid = _period_id(c, "2026.3Q")

    # 이관 기간(마감됨)에는 입고 불가
    r = c.post("/receipts", data={"date": "2026-06-30", "item_code": "NP-0001", "location_id": 1, "qty": 1, "amount": 1},
               headers={"referer": "http://testserver/receipts"})
    assert "마감된 기간" in r.text

    c.post("/receipts", data={"date": "2026-07-10", "item_code": "NP-0001 RAM : DDR4 32GB", "location_id": 1,
                              "qty": 10, "amount": 1_300_000, "vendor": "(주)케이씨티원", "voucher_no": "50015"})
    c.post("/receipts", data={"date": "2026-08-01", "item_code": "NP-0003", "location_id": 2, "qty": 5, "amount": 50_000})
    c.post("/transfers", data={"date": "2026-08-05", "item_code": "NP-0001", "location_id": 1, "qty": 1,
                               "account": "광고선전비", "note": "전시회"})
    c.post("/moves", data={"date": "2026-08-06", "item_code": "NP-0001", "from_location_id": 1, "to_location_id": 3, "qty": 2})

    # 실사: NP-0001 본사 3 + 용인 6 + 성남 2 = 11, NP-0002 미입력, NP-0003 용인 5
    r = c.get(f"/periods/{pid}/count")
    assert r.status_code == 200
    c.post(f"/periods/{pid}/count", data={"c__NP-0001__1": "3", "c__NP-0001__2": "6", "c__NP-0001__3": "2",
                                          "c__NP-0002__1": "", "c__NP-0003__2": "5"})

    r = c.get(f"/periods/{pid}")
    assert "실사 미입력 1건" in r.text

    from app import db, services
    from app.models import Period
    with db.SessionLocal() as s:
        rows = {x.code: x for x in services.ledger_rows(s, s.get(Period, pid))}
    a = rows["NP-0001"]  # 기초 10개 1,000,000 + 입고 10개 1,300,000 → 평균 115,000
    assert (a.avg_cost, a.xfer_amount, a.count_qty, a.close_amount, a.out_qty, a.out_amount) == (115000, 115000, 11, 1_265_000, 8, 920_000)
    assert rows["NP-0002"].out_qty == 3 and rows["NP-0002"].out_amount == 1_000_000
    assert rows["NP-0003"].out_qty == 0 and rows["NP-0003"].close_amount == 50_000

    # 분개장 대사
    journal = xlsx({"분개장": [["전표일자", "전표번호", "구분", "Code", "계정과목", "차변", "대변", "거래처", "Code", "적요"],
                              [date(2026, 7, 10), "50015", "차변", 153, "원재료", 1_300_000, 0, "(주)케이씨티원", None, "RAM"],
                              [date(2026, 8, 1), "50020", "차변", 153, "원재료", 40_000, 0, "11번가", None, "케이블"],
                              [date(2026, 8, 1), "50020", "대변", 251, "외상매입금", 0, 40_000, "11번가", None, "케이블"],
                              [date(2026, 10, 1), "1", "차변", 153, "원재료", 999, 0, "", None, "기간 밖"]]})
    r = c.post(f"/periods/{pid}/journal", files={"file": ("j.xlsx", journal)})
    assert "분개 3건" in r.text and "기간 밖" in r.text
    assert "확인 필요" in r.text  # 8/1 입고 50,000 vs 분개 40,000

    # 마감 → 다음 분기 기초 확정
    r = c.post(f"/periods/{pid}/close")
    assert "마감 완료" in r.text
    r = c.post("/receipts", data={"date": "2026-09-30", "item_code": "NP-0001", "location_id": 1, "qty": 1, "amount": 1},
               headers={"referer": "http://testserver/receipts"})
    assert "마감된 기간" in r.text
    c.post("/periods/next")
    with db.SessionLocal() as s:
        q4 = s.query(Period).filter_by(name="2026.4Q").one()
        opening = {x.code: (x.open_qty, x.open_amount) for x in services.ledger_rows(s, q4)}
    assert opening == {"NP-0001": (11, 1_265_000), "NP-0003": (5, 50_000)}

    # 엑셀 내보내기
    r = c.get(f"/periods/{pid}/export.xlsx")
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["수불부(2026.09)", "재고실사(2026.09)", "입고장(2026.09)", "타계정대체(2026.09)",
                             "보관처이동(2026.09)", "분개장(2026.09)"]
    ws = wb["수불부(2026.09)"]
    assert [ws.cell(3, c).value for c in (2, 17)] == ["품목코드", "기말금액"]
    assert [ws.cell(4, c).value for c in (2, 4, 6, 7, 9, 13, 14, 15, 17)] == ["NP-0001", 10, 1_000_000, 10, 1_300_000, 8, 920_000, 11, 1_265_000]
    total = next(r for r in ws.iter_rows(values_only=True) if r[1] == "합계")
    assert total[16] == 1_265_000 + 50_000
    ws = wb["재고실사(2026.09)"]
    assert [c.value for c in ws[3]][:10] == ["NP-0001", "RAM : DDR4 32GB", 11, 115000, 1_265_000, None, None, 3, 6, 2]

    # 다음 기간이 열려 있으면 마감 취소 가능
    r = c.post(f"/periods/{pid}/reopen")
    assert "마감을 취소했습니다" in r.text


def test_close_blocked_by_count_error(admin, legacy_file):
    c = admin
    c.post("/periods/import-legacy", files={"file": ("old.xlsx", legacy_file)})
    c.post("/periods/next")
    pid = _period_id(c, "2026.3Q")
    c.post(f"/periods/{pid}/count", data={"c__NP-0001__1": "99"})
    r = c.post(f"/periods/{pid}/close", headers={"referer": f"http://testserver/periods/{pid}"})
    assert "마감할 수 없습니다" in r.text and "89개 많습니다" in r.text


def test_bulk_receipt_import_is_all_or_nothing(admin, legacy_file):
    c = admin
    c.post("/periods/import-legacy", files={"file": ("old.xlsx", legacy_file)})
    c.post("/periods/next")
    header = [None, "입고일자", "월", "품목코드", "품목명", "품목 상세", "입고수량", "단가", "입고금액", "비고", "보관처"]
    bad = xlsx({"입고장": [[], header, [None, date(2026, 7, 1), "7월", "NP-0001", "", "", 2, 0, 200, "", "용인"],
                                       [None, date(2026, 7, 2), "7월", "NP-9999", "", "", 1, 0, 100]]})
    r = c.post("/receipts/import", files={"file": ("in.xlsx", bad)}, headers={"referer": "http://testserver/receipts"})
    assert "NP-9999" in r.text and "아무것도 저장되지 않음" in r.text
    good = xlsx({"입고장": [[], header, [None, date(2026, 7, 1), "7월", "NP-0001", "", "", 2, 0, 200, "", "용인"]]})
    r = c.post("/receipts/import", files={"file": ("in.xlsx", good)})
    assert "입고 1건" in r.text and "용인" in r.text


def test_login_lockout_after_repeated_failures(admin):
    c = admin
    c.post("/logout")
    for _ in range(5):
        r = c.post("/login", data={"username": "admin", "password": "wrong-pass"})
        assert "올바르지 않습니다" in r.text
    r = c.post("/login", data={"username": "admin", "password": "password123"})
    assert "잠시 막혔습니다" in r.text
    assert c.get("/", follow_redirects=False).headers["location"] == "/login"


def test_security_headers(client):
    r = client.get("/login")
    assert r.headers["x-frame-options"] == "DENY"
    assert r.headers["x-content-type-options"] == "nosniff"
