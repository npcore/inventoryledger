"""결산 엑셀 내보내기 — 기존 '재고수불 결산파일' 시트 양식을 그대로 따른다."""

import io
from datetime import timedelta
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from .ledger import to_won, totals
from .models import JournalLine, Move, Period, Receipt, Transfer
from .services import counts_by_location, ensure_locations, ledger_rows

QTY = "#,##0"
WON = "#,##0"
UNIT = "#,##0.00"
DATE = "yyyy-mm-dd"
HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
TOTAL_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)


def _header(ws: Worksheet, row: int, col: int, names: list[str]) -> None:
    for i, n in enumerate(names):
        c = ws.cell(row, col + i, n)
        c.font = Font(bold=True)
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center")
        c.border = BOX


def _row(ws: Worksheet, row: int, col: int, values: list, formats: list[str | None], fill=None) -> None:
    for i, (v, f) in enumerate(zip(values, formats)):
        c = ws.cell(row, col + i, float(v) if isinstance(v, Decimal) else v)
        c.border = BOX
        if f:
            c.number_format = f
        if fill:
            c.fill = fill
            c.font = Font(bold=True)


def _widths(ws: Worksheet, widths: dict[str, int]) -> None:
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def build_workbook(s: Session, p: Period) -> bytes:
    suffix = f"{p.end_date:%Y.%m}"
    rows = ledger_rows(s, p)
    t = totals(rows)
    locs = ensure_locations(s)
    wb = Workbook()

    # 수불부
    ws = wb.active
    ws.title = f"수불부({suffix})"
    ws["D2"], ws["Q2"] = p.start_date - timedelta(days=1), p.end_date
    ws["D2"].number_format = ws["Q2"].number_format = DATE
    cols = ["품목코드", "품목명", "기초수량", "기초단가", "기초금액", "입고수량", "입고단가", "입고금액", "평균단가",
            "대체수량", "대체금액", "출고수량", "출고금액", "기말수량", "기말단가", "기말금액"]
    fmts = [None, None, QTY, UNIT, WON, QTY, UNIT, WON, UNIT, QTY, WON, QTY, WON, QTY, UNIT, WON]
    _header(ws, 3, 2, cols)
    r = 4
    for x in rows:
        _row(ws, r, 2, [x.code, x.name, x.open_qty, x.open_cost, x.open_amount, x.in_qty, x.in_cost, x.in_amount,
                        x.avg_cost, x.xfer_qty, x.xfer_amount, x.out_qty, x.out_amount, x.count_qty,
                        x.avg_cost if x.count_qty else 0, x.close_amount], fmts)
        r += 1
    _row(ws, r, 2, ["합계", "", t.open_qty, "", t.open_amount, t.in_qty, "", t.in_amount, "", t.xfer_qty,
                    t.xfer_amount, t.out_qty, t.out_amount, t.count_qty, "", t.close_amount], fmts, TOTAL_FILL)
    ws.cell(r + 2, 13, "매출원가 PL")
    ws.cell(r + 2, 14, t.out_amount).number_format = WON
    ws.cell(r + 3, 13, "타계정으로 대체액")
    ws.cell(r + 3, 14, t.xfer_amount).number_format = WON
    ws.freeze_panes = "D4"
    _widths(ws, {"A": 2, "B": 10, "C": 48} | {get_column_letter(i): 15 for i in range(4, 18)})

    # 재고실사
    ws = wb.create_sheet(f"재고실사({suffix})")
    ws["A1"] = f"재고현황 (기준일 : {p.end_date:%Y.%m.%d})"
    ws["A1"].font = Font(bold=True, size=13)
    _header(ws, 2, 1, ["품목코드", "품목명", "재고수량", "단가", "합계", "비고", ""] + [l.name for l in locs])
    by_loc = counts_by_location(s, p)
    r = 3
    for x in rows:
        if not (x.count_qty or x.close_amount):
            continue
        _row(ws, r, 1, [x.code, x.name, x.count_qty, x.avg_cost, x.close_amount, "", ""]
             + [by_loc.get((x.code, l.id), 0) for l in locs], [None, None, QTY, UNIT, WON, None, None] + [QTY] * len(locs))
        r += 1
    loc_tot = [sum(q for (c, l), q in by_loc.items() if l == loc.id) for loc in locs]
    _row(ws, r, 1, ["합계", "", t.count_qty, "", t.close_amount, "", ""] + loc_tot,
         [None, None, QTY, None, WON, None, None] + [QTY] * len(locs), TOTAL_FILL)
    ws.freeze_panes = "A3"
    _widths(ws, {"A": 10, "B": 48, "C": 10, "D": 15, "E": 16, "F": 10, "G": 2} | {get_column_letter(8 + i): 10 for i in range(len(locs))})

    # 입고장
    ws = wb.create_sheet(f"입고장({suffix})")
    _header(ws, 2, 2, ["입고일자", "월", "품목코드", "품목명", "품목 상세", "입고수량", "단가", "입고금액", "비고",
                       "보관처", "거래처", "전표번호"])
    receipts = s.scalars(select(Receipt).where(Receipt.date.between(p.start_date, p.end_date)).order_by(Receipt.date, Receipt.id)).all()
    r = 3
    for e in receipts:
        _row(ws, r, 2, [e.date, f"{e.date.month}월", e.item_code, e.item.name, "", e.qty,
                        Decimal(e.amount) / e.qty if e.qty else 0, e.amount, e.note, e.location.name, e.vendor, e.voucher_no],
             [DATE, None, None, None, None, QTY, UNIT, WON, None, None, None, None])
        r += 1
    _row(ws, r, 2, ["합계", "", "", "", "", sum(e.qty for e in receipts), "", sum(e.amount for e in receipts), "", "", "", ""],
         [None] * 5 + [QTY, None, WON] + [None] * 4, TOTAL_FILL)
    _widths(ws, {"A": 2, "B": 12, "C": 6, "D": 10, "E": 40, "F": 10, "G": 10, "H": 14, "I": 16, "J": 16, "K": 8, "L": 24, "M": 10})

    # 타계정대체
    ws = wb.create_sheet(f"타계정대체({suffix})")
    _header(ws, 2, 2, ["대체일자", "월", "품목코드", "품목명", "품목 상세", "대체수량", "단가", "대체금액", "비고", "보관처"])
    avg = {x.code: x.avg_cost for x in rows}
    transfers = s.scalars(select(Transfer).where(Transfer.date.between(p.start_date, p.end_date)).order_by(Transfer.date, Transfer.id)).all()
    r = 3
    total = 0
    for e in transfers:
        unit = avg.get(e.item_code, Decimal(0))
        amount = to_won(e.qty * unit)
        total += amount
        _row(ws, r, 2, [e.date, f"{e.date.month}월", e.item_code, e.item.name, e.note, e.qty, unit, amount, e.account, e.location.name],
             [DATE, None, None, None, None, QTY, UNIT, WON, None, None])
        r += 1
    _row(ws, r, 2, ["합계", "", "", "", "", sum(e.qty for e in transfers), "", total, "", ""],
         [None] * 5 + [QTY, None, WON, None, None], TOTAL_FILL)
    _widths(ws, {"A": 2, "B": 12, "C": 6, "D": 10, "E": 40, "F": 16, "G": 10, "H": 12, "I": 14, "J": 14, "K": 8})

    # 보관처 이동
    moves = s.scalars(select(Move).where(Move.date.between(p.start_date, p.end_date)).order_by(Move.date, Move.id)).all()
    if moves:
        ws = wb.create_sheet(f"보관처이동({suffix})")
        _header(ws, 1, 1, ["이동일자", "품목코드", "품목명", "수량", "출발", "도착", "비고"])
        for r, e in enumerate(moves, start=2):
            _row(ws, r, 1, [e.date, e.item_code, e.item.name, e.qty, e.from_location.name, e.to_location.name, e.note],
                 [DATE, None, None, QTY, None, None, None])
        _widths(ws, {"A": 12, "B": 10, "C": 40, "D": 8, "E": 8, "F": 8, "G": 20})

    # 분개장
    lines = s.scalars(select(JournalLine).where(JournalLine.period_id == p.id).order_by(JournalLine.date, JournalLine.id)).all()
    if lines:
        ws = wb.create_sheet(f"분개장({suffix})")
        _header(ws, 1, 1, ["전표일자", "전표번호", "구분", "Code", "계정과목", "차변", "대변", "거래처", "적요"])
        for r, j in enumerate(lines, start=2):
            _row(ws, r, 1, [j.date, j.voucher_no, j.side, j.account_code, j.account_name, j.debit, j.credit, j.vendor, j.memo],
                 [DATE, None, None, None, None, WON, WON, None, None])
        _widths(ws, {"A": 12, "B": 10, "C": 6, "D": 6, "E": 10, "F": 14, "G": 14, "H": 28, "I": 36})

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
