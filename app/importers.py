"""엑셀 가져오기 (.xls / .xlsx).

기존 결산 파일의 시트 양식(헤더 이름)을 기준으로 읽는다. 헤더 행 위치나 열 순서가
조금 달라도 되도록 헤더 이름으로 열을 찾는다.
"""

import io
import re
from collections.abc import Iterator
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta

import openpyxl
import xlrd
from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from .models import ClosingBalance, Count, Item, JournalLine, Location, Period, Receipt, Transfer
from .services import BusinessError, check_editable, ensure_locations

Rows = list[list[object]]


def load_workbook(filename: str, data: bytes) -> dict[str, Rows]:
    if filename.lower().endswith(".xls"):
        wb = xlrd.open_workbook(file_contents=data)
        out = {}
        for sh in wb.sheets():
            rows = []
            for r in range(sh.nrows):
                row = []
                for c in sh.row(r):
                    if c.ctype == xlrd.XL_CELL_DATE:
                        row.append(xlrd.xldate_as_datetime(c.value, wb.datemode).date())
                    else:
                        row.append(c.value)
                rows.append(row)
            out[sh.name] = rows
        return out
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    return {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}


def text(v: object) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def to_int(v: object) -> int:
    if v in (None, ""):
        return 0
    if isinstance(v, str):
        v = v.replace(",", "").strip() or "0"
    return int(round(float(v)))


def to_date(v: object) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and 20000 < v < 80000:  # 엑셀 날짜 일련번호
        return date(1899, 12, 30) + timedelta(days=int(v))
    if isinstance(v, str) and v.strip():
        for fmt in ("%Y-%m-%d", "%Y.%m.%d", "%Y/%m/%d", "%Y%m%d"):
            try:
                return datetime.strptime(v.strip()[:10], fmt).date()
            except ValueError:
                pass
    return None


def find_header(rows: Rows, required: list[str]) -> int:
    for i, row in enumerate(rows[:30]):
        cells = {text(c) for c in row}
        if all(h in cells for h in required):
            return i
    raise BusinessError(f"헤더({', '.join(required)})가 있는 행을 찾지 못했습니다.")


def table(rows: Rows, required: list[str]) -> Iterator[tuple[int, dict[str, object]]]:
    """헤더 아래 각 행을 {헤더: 값} 으로 돌려준다. 같은 헤더가 여러 번이면 첫 열을 쓴다."""
    h = find_header(rows, required)
    cols: dict[str, int] = {}
    for i, c in enumerate(rows[h]):
        cols.setdefault(text(c), i)
    for n, row in enumerate(rows[h + 1 :], start=h + 2):
        yield n, {name: (row[i] if i < len(row) else None) for name, i in cols.items()}


ITEM_CODE = re.compile(r"^[A-Z]{2,}-\d+$")


def is_item_code(v: object) -> bool:
    return bool(ITEM_CODE.match(text(v)))


def is_placeholder_name(name: str) -> bool:
    return name in ("", "0")


@dataclass
class ImportResult:
    counts: dict[str, int] = field(default_factory=dict)
    skipped: list[str] = field(default_factory=list)

    def summary(self) -> str:
        s = ", ".join(f"{k} {v}건" for k, v in self.counts.items())
        if self.skipped:
            s += f" (건너뜀 {len(self.skipped)}건)"
        return s


def _upsert_item(s: Session, code: str, name: str) -> None:
    item = s.get(Item, code)
    if item is None:
        s.add(Item(code=code, name=name))
        s.flush()
    elif name and is_placeholder_name(item.name):
        item.name = name


def latest_suffix(sheet_names: list[str]) -> str:
    suffixes = [m.group(1) for n in sheet_names if (m := re.fullmatch(r"수불부\((\d{4}\.\d{2})\)", n.strip()))]
    if not suffixes:
        raise BusinessError("'수불부(YYYY.MM)' 시트를 찾지 못했습니다.")
    return max(suffixes)


def import_legacy(s: Session, filename: str, data: bytes, user: str) -> tuple[Period, ImportResult]:
    """기존 결산 엑셀의 최신 기간을 '기준 기간'(마감 상태)으로 가져온다.

    이 기간의 기말재고가 다음 분기의 기초재고가 된다.
    """
    if s.scalar(select(Period.id).limit(1)):
        raise BusinessError("이미 기간 데이터가 있습니다. 기준 파일 가져오기는 처음 한 번만 할 수 있습니다.")
    sheets = load_workbook(filename, data)
    suffix = latest_suffix(list(sheets))
    res = ImportResult()
    locs = {l.name: l for l in ensure_locations(s)}
    hq = next(iter(locs.values()))

    # 품목
    n = 0
    for _, r in table(sheets.get("원재료 코드표", []), ["품목코드", "품목명"]):
        code, name = text(r["품목코드"]), text(r["품목명"])
        if is_item_code(code) and not is_placeholder_name(name):
            _upsert_item(s, code, name)
            n += 1
    res.counts["품목"] = n

    # 기간 + 기말잔액
    ledger = sheets[f"수불부({suffix})"]
    h = find_header(ledger, ["품목코드", "기말수량", "기말금액"])
    dates = [d for c in ledger[h - 1] if (d := to_date(c))] if h > 0 else []
    if len(dates) < 2:
        raise BusinessError("수불부 헤더 위의 기초/기말 날짜를 찾지 못했습니다.")
    # 기초 열 위 날짜는 전기말(예: 2025-12-31)이므로 기간은 그 다음 날부터 시작한다.
    start, end = dates[0] + timedelta(days=1), dates[-1]
    p = Period(name=f"{suffix} (엑셀 이관)", start_date=start, end_date=end, closed=True, legacy=True,
               closed_at=datetime.now(), closed_by=user)
    s.add(p)
    s.flush()
    n = 0
    for _, r in table(ledger, ["품목코드", "기말수량", "기말금액"]):
        code = text(r["품목코드"])
        if not is_item_code(code):
            continue
        qty, amount = to_int(r["기말수량"]), to_int(r["기말금액"])
        if qty or amount:
            _upsert_item(s, code, text(r.get("품목명")))
            unit = r.get("기말단가")
            unit_cost = repr(float(unit)) if isinstance(unit, (int, float)) and unit else str(amount / qty if qty else 0)
            s.add(ClosingBalance(period_id=p.id, item_code=code, qty=qty, unit_cost=unit_cost, amount=amount))
            n += 1
    res.counts["기말재고"] = n

    # 보관처별 실사
    if (rows := sheets.get(f"재고실사({suffix})")) is not None:
        n = 0
        for _, r in table(rows, ["품목코드", "재고수량"]):
            code = text(r["품목코드"])
            if not is_item_code(code) or s.get(Item, code) is None:
                continue
            for name, loc in locs.items():
                if name in r and (q := to_int(r[name])):
                    s.add(Count(period_id=p.id, item_code=code, location_id=loc.id, qty=q))
                    n += 1
        res.counts["실사"] = n

    # 기간 내 입고·대체·분개 내역 (이력 조회용)
    if (rows := sheets.get(f"입고장({suffix})")) is not None:
        n = 0
        for _, r in table(rows, ["입고일자", "품목코드", "입고수량", "입고금액"]):
            code, d = text(r["품목코드"]), to_date(r["입고일자"])
            if is_item_code(code) and d:
                _upsert_item(s, code, text(r.get("품목명")))
                s.add(Receipt(date=d, item_code=code, location_id=hq.id, qty=to_int(r["입고수량"]),
                              amount=to_int(r["입고금액"]), note=text(r.get("비고")), created_by=user))
                n += 1
        res.counts["입고"] = n
    if (rows := sheets.get(f"타계정대체({suffix})")) is not None:
        n = 0
        for _, r in table(rows, ["대체일자", "품목코드", "대체수량"]):
            code, d = text(r["품목코드"]), to_date(r["대체일자"])
            if is_item_code(code) and d:
                _upsert_item(s, code, text(r.get("품목명")))
                s.add(Transfer(date=d, item_code=code, location_id=hq.id, qty=to_int(r["대체수량"]),
                               account=text(r.get("비고")) or "기타", created_by=user))
                n += 1
        res.counts["타계정대체"] = n
    if (rows := sheets.get(f"분개장({suffix})")) is not None:
        res.counts["분개"] = _add_journal(s, p, rows, res)
    s.flush()
    return p, res


def import_receipts(s: Session, filename: str, data: bytes, user: str) -> ImportResult:
    """입고장 양식(입고일자·품목코드·입고수량·입고금액 [·보관처·거래처·전표번호·비고])의 엑셀을 일괄 등록."""
    sheets = load_workbook(filename, data)
    rows = next((r for name, r in sheets.items() if "입고" in name), next(iter(sheets.values())))
    locs = {l.name: l for l in ensure_locations(s)}
    default_loc = next(iter(locs.values()))
    res = ImportResult()
    errors = []
    new = []
    for line, r in table(rows, ["입고일자", "품목코드", "입고수량", "입고금액"]):
        code, d = text(r["품목코드"]), to_date(r["입고일자"])
        if not code and not d:
            continue
        if not d or not is_item_code(code):
            res.skipped.append(f"{line}행: 날짜/품목코드 형식 오류")
            continue
        if s.get(Item, code) is None:
            errors.append(f"{line}행: 등록되지 않은 품목코드 {code}")
            continue
        loc_name = text(r.get("보관처"))
        if loc_name and loc_name not in locs:
            errors.append(f"{line}행: 알 수 없는 보관처 {loc_name}")
            continue
        try:
            check_editable(s, d)
        except BusinessError as e:
            errors.append(f"{line}행: {e}")
            continue
        qty = to_int(r["입고수량"])
        if qty <= 0:
            res.skipped.append(f"{line}행: 수량 0")
            continue
        new.append(Receipt(date=d, item_code=code, location_id=locs[loc_name].id if loc_name else default_loc.id,
                           qty=qty, amount=to_int(r["입고금액"]), vendor=text(r.get("거래처")),
                           voucher_no=text(r.get("전표번호")), note=text(r.get("비고")), created_by=user))
    if errors:
        raise BusinessError("가져오지 못했습니다 (아무것도 저장되지 않음):\n" + "\n".join(errors))
    s.add_all(new)
    res.counts["입고"] = len(new)
    return res


def import_journal(s: Session, p: Period, filename: str, data: bytes) -> ImportResult:
    """분개장을 기간에 붙인다. 기존에 올린 분개장은 교체된다."""
    sheets = load_workbook(filename, data)
    rows = next((r for name, r in sheets.items() if "분개" in name), next(iter(sheets.values())))
    s.execute(delete(JournalLine).where(JournalLine.period_id == p.id))
    res = ImportResult()
    res.counts["분개"] = _add_journal(s, p, rows, res)
    return res


def _add_journal(s: Session, p: Period, rows: Rows, res: ImportResult) -> int:
    n = 0
    for line, r in table(rows, ["전표일자", "차변", "대변"]):
        d = to_date(r["전표일자"])
        if d is None:
            continue
        if not p.start_date <= d <= p.end_date:
            res.skipped.append(f"{line}행: 기간 밖 ({d})")
            continue
        s.add(JournalLine(period_id=p.id, date=d, voucher_no=text(r.get("전표번호")), side=text(r.get("구분")),
                          account_code=text(r.get("Code")), account_name=text(r.get("계정과목")),
                          debit=to_int(r["차변"]), credit=to_int(r["대변"]), vendor=text(r.get("거래처")),
                          memo=text(r.get("적요"))))
        n += 1
    return n
